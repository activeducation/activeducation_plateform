"""
Importe les établissements d'enseignement supérieur du Togo depuis
le fichier Excel Base de Données Nationale.

Stratégie : upsert par nom (clé naturelle). Les écoles présentes à la
fois dans le fichier Excel et en base sont mises à jour ; les nouvelles
sont insérées ; les écoles orphelines (présentes en base mais plus
dans le fichier) sont désactivées (is_active=False) plutôt que
supprimées, pour préserver l'historique des inscriptions utilisateur.

Audit #11 (2026-07-30) : le script supprimait TOUTES les écoles avant
réinsertion, sans confirmation ni dry-run. Une mauvaise manipulation
ou un fichier Excel corrompu pouvait wipe la base. Corrections :
- argparse avec --dry-run (défaut) et --apply (action réelle)
- Prompt interactif "TAPEZ OUI" en plus du --apply
- Backup JSON timestampé avant toute écriture destructive
- Log clair de ce qui VA être fait avant exécution

Usage:
    cd backend
    # Voir ce qui va changer sans rien écrire :
    python -m scripts.import_schools
    # Appliquer réellement :
    python -m scripts.import_schools --apply
"""

from __future__ import annotations

import argparse
import json
import os
import re
import sys
from datetime import datetime
from typing import Any

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

EXCEL_PATH = os.getenv("EXCEL_PATH", "/tmp/ecoles.xlsx")

from openpyxl import load_workbook  # noqa: E402
from app.db.supabase_client import get_admin_supabase_client  # noqa: E402

TYPE_MAPPING: dict[str, str] = {
    "Université publique": "university",
    "Université privée": "university",
    "Université privée confessionnelle": "university",
    "Université privée internationale": "university",
    "Grande école publique": "grande_ecole",
    "Grande école publique (composante Université de Lomé)": "grande_ecole",
    "Grande école publique (santé)": "grande_ecole",
    "Grande école publique (travail social)": "grande_ecole",
    "Grande école privée": "grande_ecole",
    "Grande école privée (arts)": "grande_ecole",
    "Grande école privée (audiovisuel)": "grande_ecole",
    "Grande école privée (aéronautique/technologies)": "grande_ecole",
    "Grande école privée (banque/finance)": "grande_ecole",
    "Grande école privée (business school internationale)": "grande_ecole",
    "Grande école privée (communication)": "grande_ecole",
    "Grande école privée (design/architecture)": "grande_ecole",
    "Grande école privée (ingénierie)": "grande_ecole",
    "Grande école privée (tourisme/hôtellerie)": "grande_ecole",
    "Grande école privée (économie numérique)": "grande_ecole",
    "Grande école privée confessionnelle": "grande_ecole",
    "École inter-États d'enseignement supérieur": "grande_ecole",
    "École inter-États d'enseignement supérieur et de recherche": "grande_ecole",
    "Institut supérieur privé": "institut",
    "Institut supérieur privé (informatique)": "institut",
    "Institut supérieur privé (langues/business)": "institut",
    "Institut supérieur privé (santé)": "institut",
    "Institut supérieur privé (santé/technologie)": "institut",
    "Institut supérieur privé (sécurité sociale/santé au travail)": "institut",
    "Institut supérieur privé (technique)": "institut",
    "Institut supérieur privé (école de droit)": "institut",
    "Institut supérieur privé (Groupe BK-Université)": "institut",
    "Institut supérieur privé confessionnel": "institut",
    "Institut supérieur public (composante Université de Lomé)": "institut",
    "Institut supérieur et technique privé": "institut",
    "Institut polytechnique privé": "institut",
    "Institut polytechnique privé (BTP)": "institut",
    "Institut universitaire privé": "institut",
    "Institut universitaire privé (réseau panafricain)": "institut",
    "Faculté publique (composante Université de Lomé)": "institut",
    "Établissement privé (école internationale)": "institut",
    "Établissement privé confessionnel (théologie)": "institut",
    "Établissement public (langues)": "institut",
    "Établissement public de formation agricole": "institut",
    "Établissement supérieur privé": "institut",
    "Établissement supérieur privé (hôtellerie)": "institut",
    "Établissement supérieur privé (maritime)": "institut",
    "Établissement supérieur privé (santé)": "institut",
    "Établissement supérieur privé (scientifique)": "institut",
    "Centre inter-États de formation technique supérieure": "centre_formation",
}


def _na(value: Any) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    if s.lower() in ("non disponible", "non communiqué", "", "none", "n/a", "—", "-"):
        return None
    return s


def _na_int(value: Any) -> int | None:
    s = _na(value)
    if s is None:
        return None
    nums = re.findall(r"\d[\d\s]*\d|\d", s.replace(" ", ""))
    if nums:
        try:
            return int(nums[0])
        except (ValueError, TypeError):
            pass
    return None


def _na_float(value: Any) -> float | None:
    s = _na(value)
    if s is None:
        return None
    try:
        return float(s)
    except (ValueError, TypeError):
        return None


def _na_list(value: Any) -> list[str]:
    s = _na(value)
    if s is None:
        return []
    parts = [p.strip() for p in s.replace(";", ",").split(",")]
    return [p for p in parts if p]


def _build_description(row: tuple) -> str | None:
    parts = []
    hist = _na(row[14]) if len(row) > 14 else None
    pres = _na(row[15]) if len(row) > 15 else None
    mission = _na(row[16]) if len(row) > 16 else None
    if pres:
        parts.append(pres)
    if mission:
        parts.append(mission)
    if hist:
        parts.append(hist)
    return "\n\n".join(parts) if parts else None


def _load_sheet_as_dict(ws) -> dict[str, list]:
    """Charge une feuille en dict indexé par ID_ETABLISSEMENT."""
    result: dict[str, list] = {}
    for row in ws.iter_rows(values_only=True):
        if not row or not row[0] or not str(row[0]).startswith("TG-"):
            continue
        result[str(row[0])] = list(row)
    return result


def load_schools(wb) -> list[dict]:
    """Charge tous les établissements depuis toutes les feuilles."""
    contacts = _load_sheet_as_dict(wb["Contacts"])
    locas = _load_sheet_as_dict(wb["Localisation"])
    frais = _load_sheet_as_dict(wb["Frais"])
    diplomes = _load_sheet_as_dict(wb["Diplomes"])
    admissions = _load_sheet_as_dict(wb["Admissions"])
    stats = _load_sheet_as_dict(wb["Statistiques"])
    infra = _load_sheet_as_dict(wb["Infrastructures"])

    ws_etab = wb["Etablissements"]
    schools: list[dict] = []

    for row in ws_etab.iter_rows(values_only=True):
        if not row or not row[0] or not str(row[0]).startswith("TG-"):
            continue
        excel_id = str(row[0])
        raw_type = str(row[6]).strip() if len(row) > 6 and row[6] else ""
        mapped_type = TYPE_MAPPING.get(raw_type)
        if mapped_type is None:
            print(f"  ⚠  Type inconnu pour {excel_id}: '{raw_type}' -> fallback 'institut'")
            mapped_type = "institut"

        c = contacts.get(excel_id, [])
        l = locas.get(excel_id, [])
        f = frais.get(excel_id, [])
        d = diplomes.get(excel_id, [])
        a = admissions.get(excel_id, [])
        s = stats.get(excel_id, [])
        i = infra.get(excel_id, [])

        school = {
            "name": str(row[1]).strip() if len(row) > 1 and row[1] else "",
            "type": mapped_type,
            "city": _na(l[5]) if len(l) > 5 else None,
            "address": _na(l[7]) if len(l) > 7 else None,
            "region": _na(l[2]) if len(l) > 2 else None,
            "latitude": _na_float(l[8]) if len(l) > 8 else None,
            "longitude": _na_float(l[9]) if len(l) > 9 else None,
            "phone": _na(c[2]) if len(c) > 2 else None,
            "email": _na(c[5]) if len(c) > 5 else None,
            "website": _na(c[6]) if len(c) > 6 else None,
            "description": _build_description(row),
            "logo_url": _na(row[3]) if len(row) > 3 else None,
            "founding_year": _na_int(row[13]) if len(row) > 13 else None,
            "accreditations": _na_list(row[11]) if len(row) > 11 else [],
            "tuition_range": _na(f[2]) if len(f) > 2 else None,
            "student_count": _na_int(s[2]) if len(s) > 2 else None,
            "degrees_offered": _na(d[2]) if len(d) > 2 else None,
            "admission_info": _build_admission_info(a),
            "infrastructure": _build_infrastructure(i),
            "is_public": True,
            "is_verified": True,
            "is_active": True,
            "programs_offered": [],
        }
        schools.append(school)

    return schools


def _build_admission_info(a: list) -> str | None:
    parts = []
    for idx, label in [(2, "Niveau requis"), (3, "Série BAC"), (4, "Concours"),
                       (5, "Étude de dossier"), (6, "Entretien"), (7, "Test"),
                       (10, "Dates d'inscription"), (11, "Dates de concours")]:
        val = _na(a[idx]) if len(a) > idx else None
        if val:
            parts.append(f"{label}: {val}")
    return "\n".join(parts) if parts else None


def _build_infrastructure(i: list) -> dict:
    result = {}
    for idx, label in [(3, "bâtiments"), (4, "amphithéâtres"), (5, "salles"),
                       (8, "laboratoires"), (9, "salles_informatique"),
                       (12, "wi_fi"), (13, "internat"), (14, "restaurant"),
                       (16, "sport"), (17, "infirmerie")]:
        val = _na(i[idx]) if len(i) > idx else None
        if val and val.lower() not in ("non", "non disponible", "0"):
            result[label] = val
    summary = _na(i[2]) if len(i) > 2 else None
    if summary:
        result["synthèse"] = summary
    return result if result else {}


def load_formations(ws, school_map: dict[str, str]) -> list[dict]:
    """Charge les formations et les mappe aux UUID des écoles."""
    programs: list[dict] = []
    for row in ws.iter_rows(values_only=True):
        if not row or not row[0] or row[0] == "ID_FORMATION":
            continue
        excel_id = str(row[1]).strip() if row[1] else ""
        school_uuid = school_map.get(excel_id)
        if not school_uuid:
            continue

        filiere = _na(row[4]) if len(row) > 4 else None
        option = _na(row[5]) if len(row) > 5 else None
        niveau = _na(row[6]) if len(row) > 6 else None
        diplome = _na(row[7]) if len(row) > 7 else None
        duree = _na(row[8]) if len(row) > 8 else None
        desc = _na(row[9]) if len(row) > 9 else None

        name_parts = [p for p in [filiere, option] if p]
        name = " - ".join(name_parts) if name_parts else (diplome or "Formation")

        duration_years = None
        if duree:
            nums = re.findall(r"\d+", duree)
            if nums:
                duration_years = int(nums[0])

        programs.append({
            "school_id": school_uuid,
            "name": name,
            "description": desc or None,
            "level": niveau or None,
            "duration_years": duration_years,
            "is_active": True,
            "display_order": 0,
        })

    return programs


def backup_schools(db) -> str:
    """Sauvegarde la table schools en JSON avant toute écriture.

    Retourne le chemin du fichier de backup. Le timestamp dans le nom
    garantit qu'on garde l'historique en cas de re-run.
    """
    ts = datetime.now().strftime("%Y%m%d-%H%M%S")
    backup_dir = os.path.join(os.path.dirname(EXCEL_PATH), "backups")
    os.makedirs(backup_dir, exist_ok=True)
    backup_path = os.path.join(backup_dir, f"schools-backup-{ts}.json")
    print(f"→ Backup des écoles existantes vers {backup_path}...")
    res = db.client.table("schools").select("*").execute()
    with open(backup_path, "w", encoding="utf-8") as f:
        json.dump(res.data or [], f, ensure_ascii=False, indent=2, default=str)
    print(f"  ✓ {len(res.data or [])} écoles sauvegardées")
    return backup_path


def upsert_schools(db, schools: list[dict]) -> dict[str, str]:
    """Upsert par nom (clé naturelle).

    Pour chaque école du fichier :
    - existe en DB avec le même nom → UPDATE (tous les champs sauf id et
      created_at) + is_active=True
    - n'existe pas → INSERT

    Retourne un mapping excel_id → uuid pour réutilisation par
    load_formations().
    """
    print(f"\n→ Upsert de {len(schools)} écoles...")

    # 1. Charger toutes les écoles existantes en un seul SELECT (évite N+1)
    print("  Chargement de l'état actuel en base...")
    existing_res = db.client.table("schools").select("id, name").execute()
    existing_by_name: dict[str, str] = {
        row["name"]: row["id"] for row in (existing_res.data or [])
    }
    print(f"  {len(existing_by_name)} écoles déjà en base")

    to_insert: list[dict] = []
    to_update: list[dict] = []

    for school in schools:
        school_data = {
            "name": school["name"],
            "type": school["type"],
            "city": school["city"],
            "address": school["address"],
            "region": school["region"],
            "latitude": school["latitude"],
            "longitude": school["longitude"],
            "phone": school["phone"],
            "email": school["email"],
            "website": school["website"],
            "description": school["description"],
            "logo_url": school["logo_url"],
            "founding_year": school["founding_year"],
            "accreditations": school["accreditations"],
            "tuition_range": school["tuition_range"],
            "student_count": school["student_count"],
            "degrees_offered": school["degrees_offered"],
            "admission_info": school["admission_info"],
            "infrastructure": json.dumps(school["infrastructure"]) if school["infrastructure"] else None,
            "is_public": school["is_public"],
            "is_verified": school["is_verified"],
            "is_active": school["is_active"],
            "programs_offered": school["programs_offered"],
        }
        existing_id = existing_by_name.get(school["name"])
        if existing_id:
            school_data["id"] = existing_id
            to_update.append(school_data)
        else:
            to_insert.append(school_data)

    # 2. Insérer les nouvelles (par batch de 50 pour éviter timeout PostgREST)
    batch_size = 50
    if to_insert:
        print(f"  Insertion de {len(to_insert)} nouvelles écoles...")
        for i in range(0, len(to_insert), batch_size):
            batch = to_insert[i:i + batch_size]
            try:
                db.client.table("schools").insert(batch).execute()
            except Exception as e:
                print(f"  ✗ Erreur batch insert {i}-{i+len(batch)}: {e}")

    # 3. Mettre à jour les existantes (une par une, PostgREST .update()
    #    sans filtre ne sait pas faire de bulk sur sous-ensembles)
    if to_update:
        print(f"  Mise à jour de {len(to_update)} écoles existantes...")
        for i, school_data in enumerate(to_update):
            sid = school_data.pop("id")
            try:
                db.client.table("schools").update(school_data).eq("id", sid).execute()
            except Exception as e:
                print(f"  ✗ Erreur update '{school_data.get('name')}': {e}")
            if (i + 1) % 20 == 0:
                print(f"    {i + 1}/{len(to_update)} mises à jour")

    print(f"  ✓ {len(to_insert)} insertions, {len(to_update)} mises à jour")

    # 4. Reconstruire le mapping excel_id → uuid en re-fetchant
    final_map: dict[str, str] = {}
    wb_map = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    ws_etab_map = wb_map["Etablissements"]
    for row in ws_etab_map.iter_rows(values_only=True):
        if not row or not row[0] or not str(row[0]).startswith("TG-"):
            continue
        excel_id = str(row[0])
        name = str(row[1]).strip() if len(row) > 1 and row[1] else ""
        if name and name in existing_by_name:
            # Idempotence : si l'école existait déjà, on garde son id
            final_map[excel_id] = existing_by_name[name]
        elif name:
            try:
                fetched = (
                    db.client.table("schools").select("id").eq("name", name).limit(1).execute()
                )
                if fetched.data:
                    final_map[excel_id] = fetched.data[0]["id"]
            except Exception as e:
                print(f"  ⚠  Erreur recherche '{name}': {e}")
    wb_map.close()
    print(f"  ✓ Mapping de {len(final_map)}/{len(schools)} écoles réussi")
    return final_map


def deactivate_orphan_schools(db, kept_ids: set[str]) -> int:
    """Désactive (is_active=False) les écoles en base absentes du fichier Excel.

    Préserve l'historique (school_programs, enrollments) lié à ces écoles
    tout en les masquant des listes publiques. Pour les réactiver, un
    re-import où l'école réapparaît suffit.
    """
    print("\n→ Recherche d'écoles orphelines (présentes en base, absentes du fichier)...")
    res = (
        db.client.table("schools")
        .select("id, name")
        .eq("is_active", True)
        .execute()
    )
    all_active = res.data or []
    orphans = [row for row in all_active if row["id"] not in kept_ids]
    if not orphans:
        print("  ✓ Aucune école orpheline")
        return 0
    print(f"  {len(orphans)} écoles à désactiver :")
    for row in orphans[:10]:
        print(f"    - {row['name']}")
    if len(orphans) > 10:
        print(f"    ... et {len(orphans) - 10} autres")
    # Désactiver par lots
    batch_size = 50
    for i in range(0, len(orphans), batch_size):
        batch_ids = [o["id"] for o in orphans[i:i + batch_size]]
        try:
            db.client.table("schools").update({"is_active": False}).in_("id", batch_ids).execute()
        except Exception as e:
            print(f"  ✗ Erreur désactivation batch {i}-{i+len(batch_ids)}: {e}")
    print(f"  ✓ {len(orphans)} écoles désactivées")
    return len(orphans)


def insert_programs(db, programs: list[dict]) -> None:
    print(f"\n→ Insertion de {len(programs)} formations...")
    batch: list[dict] = []
    batch_size = 50

    for i, prog in enumerate(programs):
        batch.append(prog)
        if len(batch) >= batch_size:
            try:
                db.client.table("school_programs").insert(batch).execute()
            except Exception as e:
                print(f"  ✗ Erreur batch formations (ligne ~{i}): {e}")
            batch = []
        if (i + 1) % 100 == 0:
            print(f"  ✓ {i + 1}/{len(programs)} formations insérées")

    if batch:
        try:
            db.client.table("school_programs").insert(batch).execute()
        except Exception as e:
            print(f"  ✗ Erreur batch final formations: {e}")

    print(f"  ✓ {len(programs)} formations insérées")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Import / sync schools from the national Excel file.",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Apply changes (default: dry-run, no writes).",
    )
    parser.add_argument(
        "--excel-path",
        default=EXCEL_PATH,
        help="Path to the .xlsx file (default: $EXCEL_PATH or /tmp/ecoles.xlsx).",
    )
    parser.add_argument(
        "--yes",
        action="store_true",
        help="Skip the interactive confirmation prompt.",
    )
    args = parser.parse_args()

    excel_path = args.excel_path

    print("=" * 60)
    print("Import Base de Données Nationale — Établissements")
    print("=" * 60)
    print(f"Fichier: {excel_path}")
    print(f"Mode: {'APPLY (écriture réelle)' if args.apply else 'DRY-RUN (lecture seule)'}")

    if not os.path.exists(excel_path):
        print(f"\n✗ Fichier introuvable: {excel_path}")
        print("  Télécharge-le depuis Google Sheets au format xlsx.")
        sys.exit(1)

    print("\nOuverture du fichier Excel...")
    wb = load_workbook(excel_path, read_only=True, data_only=True)

    schools = load_schools(wb)
    print(f"\nÉcoles chargées: {len(schools)}")

    if not schools:
        print("✗ Aucune école trouvée dans le fichier.")
        wb.close()
        sys.exit(1)

    wb.close()

    print("\nAperçu des 3 premières écoles:")
    for s in schools[:3]:
        extra = []
        if s.get("region"):
            extra.append(f"region={s['region']}")
        if s.get("latitude"):
            extra.append(f"gps={s['latitude']},{s['longitude']}")
        if s.get("tuition_range"):
            extra.append("frais=✓")
        if s.get("degrees_offered"):
            extra.append("diplomes=✓")
        if s.get("admission_info"):
            extra.append("admission=✓")
        if s.get("infrastructure"):
            extra.append("infra=✓")
        print(f"  - {s['name']} ({', '.join(extra) if extra else 'base'})")

    print("\n" + "=" * 60)
    print("Connexion à Supabase (service_role)...")
    db = get_admin_supabase_client()
    print("✓ Connecté")

    # Charger l'état actuel pour comparer
    existing_res = db.client.table("schools").select("id, name, is_active").execute()
    existing_by_name = {row["name"]: row for row in (existing_res.data or [])}
    to_insert = [s for s in schools if s["name"] not in existing_by_name]
    to_update = [s for s in schools if s["name"] in existing_by_name]
    to_deactivate = [
        row for name, row in existing_by_name.items()
        if name not in {s["name"] for s in schools} and row.get("is_active", True)
    ]
    print(f"\nPlan d'action:")
    print(f"  - {len(to_insert)} nouvelles écoles à insérer")
    print(f"  - {len(to_update)} écoles existantes à mettre à jour")
    print(f"  - {len(to_deactivate)} écoles orphelines à désactiver (is_active=False)")

    if not args.apply:
        print("\n[DRY-RUN] Aucune écriture effectuée. Relancer avec --apply pour appliquer.")
        return

    # Confirmation interactive (en plus du flag --apply, ceinture-bretelle)
    if not args.yes:
        print("\n⚠️  ATTENTION : cette opération va modifier la base de données.")
        print("    Un backup JSON sera créé avant toute écriture.")
        try:
            answer = input("    Tapez exactement 'OUI' (en majuscules) pour confirmer : ")
        except EOFError:
            answer = ""
        if answer != "OUI":
            print("\n✗ Annulé.")
            sys.exit(1)
        print("  ✓ Confirmé")

    # Backup AVANT toute écriture
    backup_path = backup_schools(db)
    print(f"  (En cas de problème : restaurez depuis {backup_path})")

    school_map = upsert_schools(db, schools)
    kept_ids = set(school_map.values())
    deactivated = deactivate_orphan_schools(db, kept_ids)

    print(f"\nChargement des formations...")
    wb2 = load_workbook(excel_path, read_only=True, data_only=True)
    programs = load_formations(wb2["Formations"], school_map)
    wb2.close()
    print(f"Formations chargées: {len(programs)}")

    if programs:
        insert_programs(db, programs)

    print("\n" + "=" * 60)
    print("✓ Import terminé avec succès!")
    print(f"  {len(to_insert)} écoles créées, {len(to_update)} mises à jour, {deactivated} désactivées")
    print(f"  {len(programs)} formations importées")
    print(f"  Backup: {backup_path}")
    print("=" * 60)


if __name__ == "__main__":
    main()
