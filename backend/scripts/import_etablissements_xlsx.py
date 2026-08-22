"""Rapprochement et import d'un fichier Excel d'etablissements superieurs.

Le fichier source fournit un ROSTER d'etablissements (noms). Les autres
colonnes (frais, infrastructures, avis, URLs...) doivent etre auditees avant
tout import : un generateur peut les avoir remplies avec des valeurs
constantes qui degraderaient les donnees reelles deja en base.

Trois etapes, volontairement separees pour qu'un humain tranche au milieu :

    1. --audit      analyse la qualite du fichier (colonnes constantes,
                    cardinalite) et dit ce qui est exploitable.
    2. --rapprocher rapproche les noms du fichier avec `schools` et produit
                    un CSV a trois etats : deja_en_base / a_verifier / nouveau.
    3. --importer   insere les lignes du CSV completees a la main
                    (name, city, type obligatoires -- NOT NULL en base).

Aucune valeur n'est inventee : l'import n'ecrit que ce que le CSV contient.

Usage :
    cd backend
    python scripts/import_etablissements_xlsx.py --audit --file fichier.xlsx
    python scripts/import_etablissements_xlsx.py --rapprocher --file fichier.xlsx --out rapprochement.csv
    # ... completer rapprochement.csv a la main (city, type, decision=importer)
    python scripts/import_etablissements_xlsx.py --importer rapprochement.csv --dry-run
    python scripts/import_etablissements_xlsx.py --importer rapprochement.csv

Prerequis :
    - .env avec SUPABASE_URL et SUPABASE_SERVICE_ROLE_KEY
    - pip install openpyxl
"""

from __future__ import annotations

import argparse
import csv
import difflib
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

# Colonne portant le nom de l'etablissement, cherchee dans chaque feuille.
NAME_HEADERS = ("etablissement", "ecole", "nom de l etablissement")

# Valeurs acceptees par la contrainte CHECK de schools.type.
VALID_TYPES = ("university", "grande_ecole", "institut", "centre_formation")

CSV_FIELDS = [
    "decision",       # importer | ignorer  (rempli par l'humain)
    "statut",         # deja_en_base | a_verifier | nouveau  (calcule)
    "name",
    "city",           # NOT NULL en base -- a completer
    "type",           # NOT NULL en base -- a completer
    "region",
    "website",
    "description",
    "correspondance_base",
    "similarite",
]


# --------------------------------------------------------------------------
# Normalisation / rapprochement de noms
# --------------------------------------------------------------------------

def _strip_accents(text: str) -> str:
    text = str(text).replace("’", "'")
    decomposed = unicodedata.normalize("NFKD", text)
    return "".join(c for c in decomposed if not unicodedata.combining(c))


def normalize(name: str) -> str:
    """Cle de comparaison : sans accents, sans parentheses, sans ponctuation."""
    text = re.sub(r"\(.*?\)", " ", _strip_accents(name)).lower()
    return " ".join(re.sub(r"[^a-z0-9]+", " ", text).split())


def acronyms(name: str) -> set[str]:
    """Acronymes entre parentheses, ex: ESGIS dans '... (ESGIS)'."""
    found = re.findall(r"\(([^)]*)\)", _strip_accents(name))
    out = set()
    for raw in found:
        letters = re.sub(r"[^A-Za-z]", "", raw).upper()
        if 2 <= len(letters) <= 12:
            out.add(letters)
    return out


# --------------------------------------------------------------------------
# Lecture du fichier
# --------------------------------------------------------------------------

def load_workbook_rows(path: Path) -> dict[str, list[dict]]:
    try:
        import openpyxl
    except ImportError:
        print("ERREUR : openpyxl requis -> pip install openpyxl", file=sys.stderr)
        raise SystemExit(1)

    wb = openpyxl.load_workbook(path, data_only=True)
    sheets: dict[str, list[dict]] = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(h) if h is not None else "" for h in rows[0]]
        sheets[ws.title] = [dict(zip(headers, r)) for r in rows[1:]]
    return sheets


def extract_names(sheets: dict[str, list[dict]]) -> list[str]:
    """Recupere la liste des noms d'etablissements, sans supposer la feuille."""
    for rows in sheets.values():
        if not rows:
            continue
        for header in rows[0]:
            if normalize(header) in NAME_HEADERS:
                names = [str(r[header]).strip() for r in rows if r.get(header)]
                # Dedoublonnage sur le nom COMPLET : deux etablissements
                # distincts peuvent partager le meme libelle et ne differer
                # que par leur acronyme (ex: "... (CFP ANCILA)" vs
                # "... (ESFP-FIMAC)"), que normalize() supprime.
                seen: set[str] = set()
                unique = []
                for n in names:
                    key = " ".join(_strip_accents(n).lower().split())
                    if key and key not in seen:
                        seen.add(key)
                        unique.append(n)
                if unique:
                    return unique
    return []


# --------------------------------------------------------------------------
# 1. Audit qualite
# --------------------------------------------------------------------------

def run_audit(sheets: dict[str, list[dict]]) -> int:
    print("AUDIT QUALITE DU FICHIER")
    print("=" * 62)
    total_cols = 0
    constant_cols = 0
    for title, rows in sheets.items():
        if not rows:
            continue
        print(f"\n[{title}]  {len(rows)} lignes")
        for header in rows[0]:
            values = [r.get(header) for r in rows]
            counts = Counter("" if v is None else str(v).strip() for v in values)
            distinct = len(counts)
            total_cols += 1
            if distinct <= 1:
                constant_cols += 1
                only = next(iter(counts))
                print(f"  CONSTANTE  {header:26.26} -> {only[:48]!r}")
            elif distinct <= 3:
                top = ", ".join(f"{v[:22]!r}x{n}" for v, n in counts.most_common(3))
                print(f"  quasi-cst  {header:26.26} -> {top}")
            else:
                print(f"  variable   {header:26.26} -> {distinct} valeurs distinctes")

    print("\n" + "=" * 62)
    print(f"{constant_cols}/{total_cols} colonnes ont une valeur unique sur toutes les lignes.")
    if constant_cols:
        print("Ces colonnes ne portent aucune information par etablissement :")
        print("les importer ecraserait des donnees reelles par une valeur generique.")
    return 0


# --------------------------------------------------------------------------
# Connexion base
# --------------------------------------------------------------------------

def get_client():
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    from dotenv import load_dotenv
    load_dotenv(Path(__file__).resolve().parents[1] / ".env")

    import os
    from supabase import create_client

    url = os.environ.get("SUPABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        print("ERREUR : SUPABASE_URL et SUPABASE_SERVICE_ROLE_KEY requis (.env).",
              file=sys.stderr)
        raise SystemExit(1)
    return create_client(url, key)


def fetch_schools(client) -> list[dict]:
    rows: list[dict] = []
    page = 0
    while True:
        chunk = (
            client.table("schools")
            .select("id, name, city, type")
            .range(page * 1000, page * 1000 + 999)
            .execute()
            .data
            or []
        )
        rows.extend(chunk)
        if len(chunk) < 1000:
            return rows
        page += 1


# --------------------------------------------------------------------------
# 2. Rapprochement
# --------------------------------------------------------------------------

def run_rapprocher(names: list[str], out_path: Path, threshold: float) -> int:
    client = get_client()
    existing = fetch_schools(client)
    by_key = {normalize(s["name"]): s for s in existing}
    keys = list(by_key)

    counts: Counter = Counter()
    rows_out = []
    for name in names:
        key = normalize(name)
        if key in by_key:
            statut, match, ratio = "deja_en_base", by_key[key]["name"], 1.0
        else:
            close = difflib.get_close_matches(key, keys, n=1, cutoff=threshold)
            if not close:
                statut, match, ratio = "nouveau", "", 0.0
            else:
                match = by_key[close[0]]["name"]
                ratio = round(difflib.SequenceMatcher(None, key, close[0]).ratio(), 3)
                shared = bool(acronyms(name) & acronyms(match))
                statut = "deja_en_base" if (ratio >= 0.95 and shared) else "a_verifier"
        counts[statut] += 1
        rows_out.append({
            "decision": "" if statut == "deja_en_base" else "A_TRANCHER",
            "statut": statut,
            "name": name,
            "city": "",
            "type": "",
            "region": "",
            "website": "",
            "description": "",
            "correspondance_base": match,
            "similarite": ratio or "",
        })

    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=CSV_FIELDS, delimiter=";")
        writer.writeheader()
        writer.writerows(rows_out)

    print(f"Fichier : {len(names)} etablissements   Base : {len(existing)} etablissements")
    for statut in ("deja_en_base", "a_verifier", "nouveau"):
        print(f"  {statut:14} {counts[statut]}")
    print(f"\nCSV ecrit : {out_path}")
    print("Completer 'city' et 'type', puis mettre decision=importer sur les lignes voulues.")
    print(f"Types valides : {', '.join(VALID_TYPES)}")
    return 0


# --------------------------------------------------------------------------
# 3. Import
# --------------------------------------------------------------------------

def run_importer(csv_path: Path, dry_run: bool) -> int:
    with csv_path.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh, delimiter=";"))

    selected = [r for r in rows if (r.get("decision") or "").strip().lower() == "importer"]
    if not selected:
        print("Aucune ligne avec decision=importer. Rien a faire.")
        return 0

    client = get_client()
    existing_keys = {normalize(s["name"]) for s in fetch_schools(client)}

    errors: list[str] = []
    payloads: list[dict] = []
    for i, row in enumerate(selected, start=2):
        name = (row.get("name") or "").strip()
        city = (row.get("city") or "").strip()
        stype = (row.get("type") or "").strip()
        if not name:
            errors.append(f"ligne {i}: name vide")
            continue
        if not city:
            errors.append(f"ligne {i}: city obligatoire ({name})")
        if stype not in VALID_TYPES:
            errors.append(f"ligne {i}: type invalide {stype!r} ({name})")
        if normalize(name) in existing_keys:
            print(f"  = deja en base, ignore : {name}")
            continue

        payload = {
            "name": name,
            "city": city,
            "type": stype,
            "is_verified": False,   # fiche a documenter, non validee
            "is_active": True,
        }
        for optional in ("region", "website", "description"):
            value = (row.get(optional) or "").strip()
            if value:
                payload[optional] = value
        payloads.append(payload)

    if errors:
        print("ERREURS - rien n'a ete ecrit :", file=sys.stderr)
        for err in errors:
            print("  -", err, file=sys.stderr)
        return 1

    if not payloads:
        print("Toutes les lignes selectionnees sont deja en base.")
        return 0

    print(f"{len(payloads)} etablissement(s) a inserer :")
    for p in payloads:
        print(f"  + {p['name']}  ({p['city']}, {p['type']})")

    if dry_run:
        print("\n--dry-run : aucune ecriture.")
        return 0

    inserted = 0
    for p in payloads:
        try:
            client.table("schools").insert(p).execute()
            inserted += 1
        except Exception as exc:  # noqa: BLE001
            print(f"  ECHEC {p['name']} : {exc}", file=sys.stderr)
    print(f"\n{inserted}/{len(payloads)} insere(s).")
    return 0 if inserted == len(payloads) else 1


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Audit, rapprochement et import d'un roster d'etablissements.",
    )
    parser.add_argument("--file", type=Path, help="fichier .xlsx source")
    parser.add_argument("--audit", action="store_true", help="analyse qualite du fichier")
    parser.add_argument("--rapprocher", action="store_true",
                        help="rapproche les noms avec la table schools")
    parser.add_argument("--out", type=Path, default=Path("rapprochement_etablissements.csv"),
                        help="CSV produit par --rapprocher")
    parser.add_argument("--seuil", type=float, default=0.70,
                        help="seuil de similarite pour proposer un rapprochement")
    parser.add_argument("--importer", type=Path, metavar="CSV",
                        help="insere les lignes decision=importer du CSV")
    parser.add_argument("--dry-run", action="store_true", help="apercu sans ecriture")
    args = parser.parse_args()

    if args.importer:
        return run_importer(args.importer, args.dry_run)

    if not args.file:
        parser.error("--file requis pour --audit et --rapprocher")
    if not args.file.exists():
        print(f"ERREUR : fichier introuvable {args.file}", file=sys.stderr)
        return 1

    sheets = load_workbook_rows(args.file)

    if args.audit:
        return run_audit(sheets)

    if args.rapprocher:
        names = extract_names(sheets)
        if not names:
            print("ERREUR : aucune colonne d'etablissement trouvee dans le fichier.",
                  file=sys.stderr)
            return 1
        return run_rapprocher(names, args.out, args.seuil)

    parser.print_help()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
